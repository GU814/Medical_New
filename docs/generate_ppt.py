# -*- coding: utf-8 -*-
"""医疗问诊智能体 · 核心机制与原理 PPT 生成脚本（22 页扩充版 · 含测试数据可视化）
风格：浅蓝 + 白 + 深蓝，浅色调，去 AI 化（清爽、自然、留白充足）
依赖：python-pptx；配图由 make_charts.py 生成
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 配色 ----------
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BG         = RGBColor(0xF3, 0xF7, 0xFC)
DEEP       = RGBColor(0x0F, 0x3D, 0x6E)
BLUE       = RGBColor(0x2E, 0x7F, 0xC4)
LBLUE      = RGBColor(0xDD, 0xEA, 0xF7)
LBLUE2     = RGBColor(0xEF, 0xF5, 0xFB)
TEXT       = RGBColor(0x22, 0x33, 0x4A)
GREY       = RGBColor(0x6B, 0x7A, 0x8C)
TEAL       = RGBColor(0x1F, 0x9E, 0x8E)
LINE       = RGBColor(0xC9, 0xDA, 0xEC)
WARN       = RGBColor(0xC0, 0x6A, 0x2B)

FONT = "Microsoft YaHei"
CHARTS = r"C:\Users\lenovo\Desktop\作品集\medical_bot\docs\charts"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_cjk(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', name)
    cs = rPr.find(qn('a:cs'))
    if cs is None:
        cs = rPr.makeelement(qn('a:cs'), {}); rPr.append(cs)
    cs.set('typeface', name)


def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    return tb, tf


def set_run(p, text, size=14, color=TEXT, bold=False, italic=False, font=FONT, space=None):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    if space is not None:
        r.font.spacing = Pt(space)
    _set_cjk(r, font)
    return r


def header(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.18), fill=DEEP)
    rect(s, 0, Inches(1.18), SW, Pt(3), fill=BLUE)
    tb, tf = textbox(s, Inches(0.62), Inches(0.16), Inches(12), Inches(0.34))
    set_run(tf.paragraphs[0], kicker, size=12, color=RGBColor(0xBF, 0xD8, 0xF2), bold=True)
    tb2, tf2 = textbox(s, Inches(0.62), Inches(0.46), Inches(12.1), Inches(0.62))
    set_run(tf2.paragraphs[0], title, size=25, color=WHITE, bold=True)


def footer(s, idx, total=22):
    tb, tf = textbox(s, Inches(0.62), Inches(7.04), Inches(9), Inches(0.34), anchor=MSO_ANCHOR.MIDDLE)
    set_run(tf.paragraphs[0], "医疗问诊智能体 · 核心机制与原理", size=9, color=GREY)
    tb2, tf2 = textbox(s, Inches(11.0), Inches(7.04), Inches(1.7), Inches(0.34), anchor=MSO_ANCHOR.MIDDLE)
    p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.RIGHT
    set_run(p2, f"{idx:02d} / {total}", size=10, color=BLUE, bold=True)


def bullets(s, x, y, w, h, items, size=15, gap=8, color=TEXT, lead_color=BLUE):
    tb, tf = textbox(s, x, y, w, h)
    first = True
    for it in items:
        text, lvl = (it if isinstance(it, tuple) else (it, 0))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap); p.space_before = Pt(0); p.level = lvl
        if lvl == 0:
            set_run(p, "●  ", size=size, color=lead_color, bold=True)
            set_run(p, text, size=size, color=color)
        else:
            set_run(p, "    –  ", size=size-1, color=GREY, bold=True)
            set_run(p, text, size=size-1, color=color)
    return tb


def callout(s, x, y, w, h, title, body, fill=LBLUE, accent=BLUE, tcolor=DEEP, body_size=11.5):
    rect(s, x, y, w, h, fill=fill, line=accent, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tb, tf = textbox(s, x+Pt(12), y+Pt(8), w-Pt(24), h-Pt(14))
    set_run(tf.paragraphs[0], title, size=13.5, color=tcolor, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(2)
    set_run(p2, body, size=body_size, color=TEXT)
    return tb


def add_table(s, x, y, w, data, col_w=None, row_h=Pt(30), font_size=11, header=True):
    rows, cols = len(data), len(data[0])
    total_h = row_h * rows
    gfx = s.shapes.add_table(rows, cols, x, y, w, total_h)
    tbl = gfx.table
    tbl.first_row = header; tbl.horz_banding = False
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = Pt(7); cell.margin_right = Pt(7)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            if i == 0 and header:
                cell.fill.solid(); cell.fill.fore_color.rgb = DEEP
                set_run(p, str(val), size=font_size, color=WHITE, bold=True)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if (i % 2 == 1) else LBLUE2
                set_run(p, str(val), size=font_size-0.5, color=TEXT, bold=(j == 0))
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Emu(int(w * cw / tot))
    return tbl


def pic(s, name, x, y, width):
    path = os.path.join(CHARTS, name)
    return s.shapes.add_picture(path, x, y, width=width)


def flow_row(s, items, top, box_w, box_h, area_l=Inches(0.55), area_r=Inches(12.78),
             fill=LBLUE, line=BLUE, tcolor=DEEP, size=11.5, arrow=BLUE):
    """横向流程条：圆角框 + 右箭头，文本居中，支持 \n 两行（首行加粗）"""
    n = len(items)
    area_w = area_r - area_l
    gap = (area_w - box_w*n) / (n-1)
    for i, txt in enumerate(items):
        x = area_l + i*(box_w+gap)
        rect(s, x, top, box_w, box_h, fill=fill, line=line, line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tb, tf = textbox(s, x, top, box_w, box_h, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        for j, ln in enumerate(txt.split("\n")):
            pp = p if j == 0 else tf.add_paragraph(); pp.alignment = PP_ALIGN.CENTER
            set_run(pp, ln, size=size, color=tcolor, bold=(j == 0))
        if i < n-1:
            ax = x + box_w + gap/2 - Inches(0.12)
            rect(s, ax, top+box_h/2-Inches(0.14), Inches(0.24), Inches(0.28),
                 fill=arrow, shape=MSO_SHAPE.RIGHT_ARROW)


# =====================================================================
# 1. 封面
# =====================================================================
s = add_slide()
rect(s, 0, 0, SW, SH, fill=BG)
rect(s, 0, 0, Inches(0.42), SH, fill=DEEP)
rect(s, 0, 0, SW, Inches(0.16), fill=BLUE)
tb, tf = textbox(s, Inches(1.0), Inches(2.05), Inches(11.3), Inches(1.6))
set_run(tf.paragraphs[0], "医疗问诊智能体", size=46, color=DEEP, bold=True, space=1)
tb2, tf2 = textbox(s, Inches(1.02), Inches(3.15), Inches(11.3), Inches(1.0))
set_run(tf2.paragraphs[0], "核心机制与原理", size=28, color=BLUE, bold=True)
p2b = tf2.add_paragraph()
set_run(p2b, "RAG 检索增强生成  ·  Agent 自主推理  ·  防幻觉机制", size=16, color=GREY)
rect(s, Inches(1.02), Inches(4.35), Inches(4.6), Pt(2.2), fill=LINE)
tb3, tf3 = textbox(s, Inches(1.02), Inches(4.55), Inches(11), Inches(1.4))
set_run(tf3.paragraphs[0], "基于本地大模型（Ollama）的医学健康咨询系统", size=15, color=TEXT)
p3b = tf3.add_paragraph()
set_run(p3b, "从知识检索、受控推理到可信回答的工程化拆解（含实测数据）", size=15, color=TEXT)
tb4, tf4 = textbox(s, Inches(1.02), Inches(6.55), Inches(11), Inches(0.5))
set_run(tf4.paragraphs[0], "技术架构与实现解析  |  知识库 RAG · ReAct Agent · 防幻觉护栏", size=11, color=GREY)


# =====================================================================
# 2. 项目概览与内容导航
# =====================================================================
s = add_slide(); header(s, "OVERVIEW · 项目概览", "项目定位与技术栈")
bullets(s, Inches(0.62), Inches(1.5), Inches(7.0), Inches(4.6), [
    "项目定位：面向用药 / 症状 / 疾病 / 日常护理的医学健康科普智能体，可本地离线部署运行。",
    "接入形态：FastAPI 后端同时服务 Taro 小程序、桌面 Web 端，并通过 SSE 流式输出。",
    "模型配置：双 LLM —— 指令模型（qwen2.5:7b-instruct，低延迟问答）与推理模型（deepseek-r1，深度分析）；向量模型 nomic-embed-text（768 维）。",
    "知识底座：ChromaDB 向量库（medical_knowledge + consultation_memory 集合）+ SQLite（会话 / 就诊记录密文）。",
    "安全合规：AES-256-GCM 字段级加密、JWT 鉴权、紧急症状关键词前置就医提醒。",
], size=14.5, gap=10)
callout(s, Inches(7.95), Inches(1.55), Inches(4.75), Inches(1.7), "本 PPT 主线",
        "RAG 搭建 → Agent 搭建 → 防幻觉机制，\n三层能力递进：让回答“召得全、\n推得清、不编造”。", fill=LBLUE2, accent=BLUE)
callout(s, Inches(7.95), Inches(3.45), Inches(4.75), Inches(2.7), "内容导航",
        "① 背景需求与整体架构\n② 实测：模型性能与测试体系\n③ RAG：原理 · 构建 · 运维\n④ Agent：动机 · 循环 · 容错\n⑤ 防幻觉：总览 + 三道防线\n⑥ 多模态 · 演进路线 · 总结",
        fill=LBLUE2, accent=TEAL)
footer(s, 2)


# =====================================================================
# 3. 背景与需求分析（新增）
# =====================================================================
s = add_slide(); header(s, "BACKGROUND · 需求分析", "医疗问答场景的三大挑战")
bg_cards = [
    ("幻觉风险", "通用 LLM 会“一本正经地编造”剂量、禁忌与诊疗建议；医疗场景容错率极低，回答必须“有出处、可验证”，不能依赖模型的参数记忆。", WARN),
    ("隐私合规", "问诊记录、检查报告属敏感个人健康信息，上传云端大模型存在数据出域风险；需要本地离线部署、落库加密、数据不出域。", BLUE),
    ("算力成本", "部署环境为消费级 GPU / 纯 CPU，只能承载 7B 级模型；必须在有限算力下同时保证低延迟、可控性与回答质量。", TEAL),
]
cw, ch = Inches(3.90), Inches(3.30)
gx, gy = Inches(0.62), Inches(1.60)
spx = Inches(0.20)
for i, (t, b, col) in enumerate(bg_cards):
    x = gx + i*(cw+spx)
    rect(s, x, gy, cw, Inches(0.62), fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tbx, tfx = textbox(s, x, gy, cw, Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
    px = tfx.paragraphs[0]; px.alignment = PP_ALIGN.CENTER
    set_run(px, t, size=16, color=WHITE, bold=True)
    rect(s, x, gy+Inches(0.62), cw, ch-Inches(0.62), fill=LBLUE2, line=LINE, line_w=1.0)
    tb2, tf2 = textbox(s, x+Pt(14), gy+Inches(0.62)+Pt(10), cw-Pt(28), ch-Inches(0.62)-Pt(18))
    set_run(tf2.paragraphs[0], b, size=12.5, color=TEXT)
callout(s, Inches(0.62), Inches(5.35), Inches(12.1), Inches(1.15), "应对主线",
        "RAG 知识锚定（召得全：回答基于本地医学文档而非臆测）→  ReAct 受控推理（推得清：多步查证、每步可观测）→  三层防幻觉（不编造：检索 / 引用 / 工程层层设闸）。",
        fill=LBLUE, accent=BLUE, body_size=12)
footer(s, 3)


# =====================================================================
# 4. 整体架构
# =====================================================================
s = add_slide(); header(s, "ARCHITECTURE · 整体架构", "三层架构与数据流向")
band_top = Inches(1.55); band_h = Inches(1.32); gap_b = Inches(0.22)
layers = [
    ("接入层", ["小程序（Taro）/ 桌面 Web", "SSE 流式通道 · JWT 鉴权", "微信生态：订阅消息 / 小程序码"], BLUE),
    ("服务层", ["意图路由（FastAPI）", "直接问答 / 5 阶段问诊 / 报告生成", "ReAct 编排 · Ollama OpenAI 兼容调用"], DEEP),
    ("存储与知识层", ["ChromaDB：medical_knowledge + consultation_memory", "SQLite：会话 / 就诊记录（AES 加密密文）", "本地知识文档 .md / .txt / .pdf"], TEAL),
]
for i, (lab, items, col) in enumerate(layers):
    top = band_top + i * (band_h + gap_b)
    rect(s, Inches(0.62), top, Inches(2.5), band_h, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tb, tf = textbox(s, Inches(0.62), top, Inches(2.5), band_h, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    set_run(p, lab, size=17, color=WHITE, bold=True)
    rect(s, Inches(3.3), top, Inches(9.4), band_h, fill=LBLUE2, line=LINE, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tb2, tf2 = textbox(s, Inches(3.55), top+Pt(8), Inches(9.0), band_h-Pt(14))
    first = True
    for it in items:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph(); first = False
        p.space_after = Pt(2)
        set_run(p, "▸ ", size=12.5, color=col, bold=True)
        set_run(p, it, size=12.5, color=TEXT)
arrow_top = band_top + 3*(band_h+gap_b) + Inches(0.08)
rect(s, Inches(0.62), arrow_top, Inches(12.1), Inches(0.5), fill=DEEP, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tb, tf = textbox(s, Inches(0.62), arrow_top, Inches(12.1), Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
set_run(p, "数据流向：用户问题  →  意图识别  →  检索 / 工具调用  →  证据整合  →  流式可信回答",
        size=13.5, color=WHITE, bold=True)
footer(s, 4)


# =====================================================================
# 5. 端到端请求生命周期（新增）
# =====================================================================
s = add_slide(); header(s, "PIPELINE · 端到端流程", "一次问诊问答的完整旅程")
flow_row(s, ["用户提问\n小程序 / Web", "鉴权路由\nJWT · 意图识别", "并行预检索\nkb + memory",
             "ReAct 循环\n思考-行动-观察", "句子级溯源\n引用池 + [n] 绑定", "流式落库\nSSE 下发 · AES 加密"],
         top=Inches(1.85), box_w=Inches(1.95), box_h=Inches(1.35))
bullets(s, Inches(0.62), Inches(3.55), Inches(12.1), Inches(2.6), [
    "流式体验：token 级 SSE 推送，首字延迟主要取决于并行预检索与 Planner 首步，后续逐句下发、边生成边展示。",
    "事件协议：thinking（思维链）/ step（推理时间线）/ token（正文）/ citation（引用）/ done（收尾）分类下发，前端按需渲染；老前端自动忽略新事件。",
    "双通道记忆：回答完成后写入 consultation_memory 向量集合，供后续问诊语义召回，形成“越用越懂用户”的记忆闭环。",
    "落库安全：会话历史 / 报告 / 就诊记录经 AES-256-GCM 加密后写入 SQLite，密钥本地保管，全程数据不出域。",
], size=13, gap=9)
callout(s, Inches(0.62), Inches(6.15), Inches(12.1), Inches(0.62), "设计要点",
        "每一跳都可观测（SSE 事件透出）、可降级（超时 / 异常均有兜底）、可追溯（引用编号直达知识片段）。",
        fill=LBLUE, accent=BLUE, body_size=11.5)
footer(s, 5)


# =====================================================================
# 6. 实测：模型性能（图 1 + 对比表）
# =====================================================================
s = add_slide(); header(s, "BENCHMARK · 实测数据", "模型性能：切换前后单轮响应耗时")
pic(s, "1_model_latency.png", Inches(0.55), Inches(1.5), Inches(7.5))
callout(s, Inches(0.62), Inches(5.15), Inches(7.35), Inches(1.5), "关键结论",
        "切换 qwen2.5:7b-instruct 后：问诊单轮从均值 44.3s 降至 2.5~15s；"
        "报告 / 医院分析从 665s 降至 43s（约 15.5 倍提速），且不再输出不可控的长思维链。",
        fill=LBLUE, accent=BLUE)
add_table(s, Inches(8.35), Inches(1.55), Inches(4.4), [
    ["指标", "deepseek-r1:8b", "qwen2.5:7b"],
    ["单轮问诊耗时", "25.7~64.8s（均值 44.3）", "2.5~15s"],
    ["报告 / 医院分析", "665s", "43s"],
    ["思维链", "长思维链、不可控", "指令模型、无冗长链"],
    ["主链路角色", "深度分析报告", "问诊 / 直接问答"],
], col_w=[1.0, 1.5, 1.25], row_h=Pt(38), font_size=10.5)
tb, tf = textbox(s, Inches(8.35), Inches(5.55), Inches(4.4), Inches(1.2))
set_run(tf.paragraphs[0], "数据来源", size=11, color=DEEP, bold=True)
p = tf.add_paragraph()
set_run(p, "_bench ~ _bench4 共 14 次实测采样；qwen2.5 数据为 2026-09-22 切换实测。",
        size=10, color=GREY)
footer(s, 6)


# =====================================================================
# 7. 测试体系与回归保障（新增）
# =====================================================================
s = add_slide(); header(s, "TESTING · 质量保障", "77 条用例的回归测试体系")
bullets(s, Inches(0.62), Inches(1.55), Inches(5.6), Inches(4.5), [
    "用例规模：11 个测试文件、77 条用例全部通过，单次回归仅 10.77s —— 每次改动都能快速验证、放心重构。",
    "分层覆盖：单元级（主题切分 / 句子溯源 / 容错解析）→ 流程级（ReAct 循环 / 问诊阶段推进）→ 接口级（微信 API / 多模态）。",
    "冒烟保障：test_react_smoke 端到端跑通“提问 → 检索 → 溯源回答”主链路，防止主流程静默退化。",
    "防幻觉专项：report_visit_date 锁死日期铁律、context_retention 验证多轮上下文保持，回归即护栏。",
], size=12.5, gap=10)
add_table(s, Inches(6.55), Inches(1.55), Inches(6.2), [
    ["测试模块", "用例数", "覆盖要点"],
    ["context_retention", "21", "多轮上下文保持"],
    ["react_cycle", "17", "ReAct 循环与降级"],
    ["multimodal", "10", "语音 / 图片识别"],
    ["chunking", "10", "主题切分质量"],
    ["report_visit_date", "7", "报告日期铁律"],
    ["thinking", "7", "思维链透出"],
    ["wx_api / location", "3 + 2", "微信 API / 位置订阅"],
], col_w=[1.5, 0.7, 1.6], row_h=Pt(34), font_size=10.5)
callout(s, Inches(0.62), Inches(6.05), Inches(12.1), Inches(0.72), "实测结论",
        "pytest 全量回归 77 passed / 10.77s，覆盖 8 大模块 —— 机制改动（切分策略、预算参数、引用规则）均有对应测试守住行为基线。",
        fill=LBLUE2, accent=TEAL, body_size=11.5)
footer(s, 7)


# =====================================================================
# 8. RAG 技术原理与选型（新增）
# =====================================================================
s = add_slide(); header(s, "RAG · 原理与选型", "为什么是检索增强生成")
bullets(s, Inches(0.62), Inches(1.5), Inches(5.6), Inches(4.6), [
    "什么是 RAG：检索（Retrieval）+ 增强（Augmented）+ 生成（Generation）—— 先查资料再作答，让模型“开卷考试”而非凭记忆答题。",
    "工作方式：用户问题向量化 → 在医学知识库中语义检索 → 把命中片段作为证据上下文注入提示词 → 模型基于证据生成并标注引用。",
    "医疗适配性：医学知识更新快、容错低，RAG 改文档即生效，且每句回答可绑定出处，天然契合“有出处、可追溯”的要求。",
], size=13, gap=10)
add_table(s, Inches(6.55), Inches(1.55), Inches(6.2), [
    ["方案", "知识更新", "可溯源", "幻觉风险", "本地成本"],
    ["纯 LLM 参数记忆", "需重新训练", "无", "高", "低"],
    ["微调（QLoRA）", "需重新训练", "无", "中", "中"],
    ["RAG（本项目）", "改文档即生效", "句子级 [n]", "低", "低"],
], col_w=[1.5, 1.1, 1.0, 0.9, 0.9], row_h=Pt(36), font_size=10.5)
callout(s, Inches(0.62), Inches(6.0), Inches(12.1), Inches(0.75), "选型结论",
        "RAG 以最低的本地算力成本同时解决“知识可更新”与“回答可溯源”两大刚需；微调管线（training/）作为后续增强手段保留，与 RAG 互补而非替代。",
        fill=LBLUE, accent=BLUE, body_size=11.5)
footer(s, 8)


# =====================================================================
# 9. RAG：知识库构建流程
# =====================================================================
s = add_slide(); header(s, "RAG · 知识库构建", "从文档到向量：索引管线")
flow_row(s, ["原始文档\n.md / .txt / .pdf", "按 ## 标题\n主题切分", "nomic-embed-text\n向量化 768 维",
             "写入 ChromaDB\ncosine 持久化", "语义检索服务\n复用问诊 / 直接问答"],
         top=Inches(1.85), box_w=Inches(2.05), box_h=Inches(1.5),
         area_l=Inches(0.7), area_r=Inches(12.6), size=12.5)
bullets(s, Inches(0.62), Inches(3.75), Inches(12.1), Inches(3.0), [
    "入库 id 约定：每个片段主键 = {文件名}_chunk_{i}，并同时写入 metadata.doc_id，保证检索端即使 include 不支持 ids 也能稳定溯源（_derive_doc_id 反推）。",
    "元数据：source（来源文件）、section（所属 ## 主题）、chunk_index（片段序号）、part / total_parts（超长块二次切分标记）。",
    "短块过滤：CHUNK_MIN_INDEX_LEN=40，过滤只剩标题行的碎块 —— 实测 14 字标题对“对乙酰氨基酚”问题虚高到 0.77，会把真正含药名的片段挤出 Top-K。",
    "检索兜底：ChromaDB 为空时自动回退本地 TF-IDF 索引（_search_local_index），Ollama 不可用时仍可提供基础检索。",
], size=13.5, gap=9)
footer(s, 9)


# =====================================================================
# 10. RAG 关键模块 + 分块质量实测（图 2）
# =====================================================================
s = add_slide(); header(s, "RAG · 关键模块与实测", "切分、向量化、检索与分块质量实测")
bullets(s, Inches(0.62), Inches(1.55), Inches(5.55), Inches(4.4), [
    "主题切分 CHUNK_STRATEGY=heading：一个 ## 标题 = 一个块，子标题与正文不跨块。",
    "向量化兼容：OpenAIEmbeddingFunction 兼容 Ollama（nomic-embed-text，768 维），含多级 fallback。",
    "检索阈值：search() 用 cosine distance，RELEVANCE_THRESHOLD=0.6，先取 2×top_k 再过滤截断。",
    "溯源去重：include 不支持 \"ids\"，用 _derive_doc_id 从 metadata 反推主键，存量集合无需重导。",
], size=13, gap=10)
pic(s, "2_chunk_quality.png", Inches(6.35), Inches(1.55), Inches(6.45))
callout(s, Inches(0.62), Inches(5.5), Inches(12.1), Inches(1.15), "实测结论",
        "11 篇文档按 ## 主题切分得 104 块（最大 1337 字），跨主题块 0 个；旧装箱策略 117 块中 114 块（97.4%）跨多标题。"
        "配合 CHUNK_MIN_INDEX_LEN=40 短块过滤，检索结果主题纯度与对症性显著提升。",
        fill=LBLUE2, accent=TEAL, body_size=11.5)
footer(s, 10)


# =====================================================================
# 11. 知识库内容与运维（新增）
# =====================================================================
s = add_slide(); header(s, "RAG · 内容与运维", "知识库构成与增量维护")
bullets(s, Inches(0.62), Inches(1.55), Inches(5.6), Inches(4.6), [
    "知识构成：11 篇本地医学文档（.md / .txt / .pdf），覆盖用药 / 症状 / 疾病 / 日常护理，按 ## 主题切分为 104 块。",
    "双集合隔离：medical_knowledge（公共医学知识）与 consultation_memory（个人问诊记忆）物理分库，用途与权限分明。",
    "人工校验：export_knowledge_chunks.py 一键导出 docs/knowledge-chunks.md，逐块人工抽查切分质量与主题纯度。",
], size=13, gap=10)
ops = [
    ("全量初始化", "init_knowledge.py / init_knowledge_local.py：清空并重建全部索引，用于首次部署或策略切换。", BLUE),
    ("增量更新", "update_knowledge.py：只重建变更文档的块，避免全量重嵌入，知识维护分钟级完成。", TEAL),
    ("全量重建", "reindex_knowledge.py：切分策略 / 向量模型变更后整体重建，保证索引与代码逻辑一致。", DEEP),
]
oy = Inches(1.55)
for i, (t, b, col) in enumerate(ops):
    y = oy + i*Inches(1.52)
    callout(s, Inches(6.55), y, Inches(6.2), Inches(1.38), t, b, fill=LBLUE2, accent=col, body_size=11.5)
callout(s, Inches(0.62), Inches(6.25), Inches(12.1), Inches(0.62), "运维原则",
        "知识更新 = 改文档 + 跑脚本，无需动代码、无需重训模型 —— RAG 架构下知识运营的成本趋近于零。",
        fill=LBLUE, accent=BLUE, body_size=11.5)
footer(s, 11)


# =====================================================================
# 12. Agent 设计动机与协议选型（新增）
# =====================================================================
s = add_slide(); header(s, "AGENT · 设计动机", "为什么是 ReAct + 结构化 JSON 协议")
bullets(s, Inches(0.62), Inches(1.5), Inches(5.6), Inches(4.6), [
    "单轮问答的局限：真实问诊常需多步查证 —— 先查知识库、再调就诊史、最后综合作答，单次“检索-生成”无法完成。",
    "小模型约束：7B 指令模型守不住自由文本 ReAct 格式，Thought/Action 混排解析失败率高，必须结构化。",
    "设计原则：工具白名单最小化（5 个）、每步可观测（SSE 透出）、任何异常都有降级出口。",
], size=13, gap=10)
add_table(s, Inches(6.55), Inches(1.55), Inches(6.2), [
    ["协议方案", "7B 稳定性", "容错能力", "结论"],
    ["自由文本 ReAct", "差：格式漂移", "弱：解析即失败", "弃用"],
    ["Function Calling", "依赖专项能力", "一般", "未采用"],
    ["JSON 协议+三级容错\n（本项目）", "稳：schema 约束", "强：可恢复", "采用"],
], col_w=[1.6, 1.2, 1.2, 0.8], row_h=Pt(38), font_size=10.5)
callout(s, Inches(0.62), Inches(6.0), Inches(12.1), Inches(0.75), "核心权衡",
        "不为 Agent 引入更大模型：Planner 与终答复用同一 qwen2.5:7b，靠“结构化协议 + 白名单 + 双闸预算”把 7B 小模型约束到可控、可恢复、可解释。",
        fill=LBLUE, accent=BLUE, body_size=11.5)
footer(s, 12)


# =====================================================================
# 13. Agent：ReAct 受控推理框架
# =====================================================================
s = add_slide(); header(s, "AGENT · 推理框架", "ReAct 受控循环：思考—行动—观察")
flow_row(s, ["用户问题", "第 0 步\n意图判定", "并行预检索\nkb+memory", "Planner\nThought/Action", "工具执行\nObservation", "终答\n句子级溯源"],
         top=Inches(1.85), box_w=Inches(1.95), box_h=Inches(1.35))
bullets(s, Inches(0.62), Inches(3.5), Inches(12.1), Inches(3.2), [
    "循环形态：Thought（Planner 思考）→ Action（工具白名单动作）→ Observation（工具真实返回值，模型不可编造）→ … → Final（终答）。",
    "采用结构化 JSON 协议而非自由文本：7B 小模型守不住自由文本 ReAct，JSON 协议 + 三级容错解析才稳。",
    "模型复用：Planner 与终答统一复用 CONSULT_MODEL_NAME（qwen2.5:7b），不新增模型标识、不额外占用显存。",
    "可解释性：完整推理步骤通过 SSE step 事件下发，前端 ReasoningSteps 组件以可折叠时间线呈现；老前端不识别该事件会自动忽略。",
    "预检索并行：kb_search 与 memory_search 用 asyncio.gather 同时发起，兼顾“零延迟起步”与“工具可被真实调用”。",
], size=13.5, gap=9)
footer(s, 13)


# =====================================================================
# 14. Agent 工具与编排 + ReAct 预算实测（图 3）
# =====================================================================
s = add_slide(); header(s, "AGENT · 工具与编排实测", "工具白名单、双闸控制与时间预算实测")
bullets(s, Inches(0.62), Inches(1.55), Inches(5.6), Inches(4.5), [
    "工具白名单：kb_search / memory_search / patient_history / finish / ask_user；白名单外一律判非法并降级。",
    "双闸控制：步数预算 REACT_MAX_STEPS=3 + 时间预算 REACT_BUDGET_MS=8000ms，超限立即强制作答。",
    "工具超时 REACT_TOOL_TIMEOUT_MS=4000ms：线程外执行 + wait_for，超时按空观察降级，不中断整轮。",
    "健壮性：Planner 连续 2 次非法输出回退单次 LLM 流程；JSON 三级容错（直解 / 去围栏 / 括号回溯）。",
    "实测依据：Planner 冷启动 5.7s，预热后降至 1.3s（qwen2.5:7b），故单次超时上限取预算 0.75 倍即 6s。",
], size=12.5, gap=9)
pic(s, "3_react_budget.png", Inches(6.4), Inches(1.6), Inches(6.35))
callout(s, Inches(6.4), Inches(4.5), Inches(6.35), Inches(2.05), "工具执行约定",
        "Observation 只来自工具真实返回值，禁止模型编造；patient_history 仅回显摘要字段；"
        "action 参数脱敏后下发（超 60 字截断）；任何工具异常收敛为 ok=False 由 loop 层降级。",
        fill=LBLUE2, accent=TEAL, body_size=11.5)
footer(s, 14)


# =====================================================================
# 15. 容错与降级链（新增）
# =====================================================================
s = add_slide(); header(s, "AGENT · 容错设计", "四级降级链：任何异常都有出口")
deg = [
    ("① 解析级容错", "Planner 输出先直解 JSON；失败则去 Markdown 围栏重试；再失败按括号回溯截取最外层 {} —— 三级容错吃掉 7B 模型的格式抖动。", BLUE),
    ("② 行为级回退", "连续 2 次非法输出（非 JSON / 白名单外动作）判定 Planner 失控，整轮回退为单次 LLM 直接问答，保证用户一定拿到回答。", TEAL),
    ("③ 工具级降级", "工具异常统一收敛为 ok=False 的空观察（Observation），循环继续而非崩溃；超时 4s 同样按空观察处理。", DEEP),
    ("④ 预算级强制作答", "步数 >3 或总耗时 >8s 触发强制 finish：用已收集证据立即生成终答；零证据则返回固定兜底文案，不推测、不编造。", WARN),
]
cw, ch = Inches(5.95), Inches(2.35)
gx, gy = Inches(0.62), Inches(1.55)
spx, spy = Inches(0.2), Inches(0.22)
for i, (t, b, col) in enumerate(deg):
    r, c = divmod(i, 2)
    x = gx + c*(cw+spx); y = gy + r*(ch+spy)
    callout(s, x, y, cw, ch, t, b, fill=LBLUE2, accent=col)
footer(s, 15)


# =====================================================================
# 16. 防幻觉体系总览（新增）
# =====================================================================
s = add_slide(); header(s, "ANTI-HALLUCINATION · 体系总览", "三道防线：从检索到工程的全链路设闸")
defense = [
    ("检索侧", "召得全、滤得准", ["宽召回绕过阈值直查", "短语加成 strong_hit 置顶", "证据门槛 0.50 过滤", "强命中豁免防误杀"], BLUE),
    ("引用侧", "有出处、可追溯", ["引用池只用 kind=kb", "doc_id 去重 + 重排", "句子级溯源 [n] 绑定", "回校验剥离虚构编号"], DEEP),
    ("工程侧", "不编造、不出域", ["AES-256-GCM 落库加密", "日期铁律禁 LLM 生成", "thinking / step 透明透出", "紧急关键词前置提醒"], TEAL),
]
cw, ch = Inches(3.90), Inches(4.35)
gx, gy = Inches(0.62), Inches(1.55)
spx = Inches(0.20)
for i, (t, sub, items, col) in enumerate(defense):
    x = gx + i*(cw+spx)
    rect(s, x, gy, cw, Inches(0.88), fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tbx, tfx = textbox(s, x, gy, cw, Inches(0.88), anchor=MSO_ANCHOR.MIDDLE)
    px = tfx.paragraphs[0]; px.alignment = PP_ALIGN.CENTER
    set_run(px, t, size=16, color=WHITE, bold=True)
    psub = tfx.add_paragraph(); psub.alignment = PP_ALIGN.CENTER
    set_run(psub, sub, size=11, color=RGBColor(0xBF, 0xD8, 0xF2))
    rect(s, x, gy+Inches(0.88), cw, ch-Inches(0.88), fill=LBLUE2, line=LINE, line_w=1.0)
    tb2, tf2 = textbox(s, x+Pt(14), gy+Inches(0.88)+Pt(12), cw-Pt(28), ch-Inches(0.88)-Pt(20))
    first = True
    for it in items:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph(); first = False
        p.space_after = Pt(8)
        set_run(p, "▸ ", size=12.5, color=col, bold=True)
        set_run(p, it, size=12.5, color=TEXT)
rect(s, Inches(0.62), Inches(6.25), Inches(12.1), Inches(0.5), fill=DEEP, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tb, tf = textbox(s, Inches(0.62), Inches(6.25), Inches(12.1), Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
set_run(p, "设计目标：医学回答“有出处、可追溯、不编造” —— 任何一环失效都有下一环兜底",
        size=13.5, color=WHITE, bold=True)
footer(s, 16)


# =====================================================================
# 17. 防幻觉（一）检索侧 + 评分实测（图 4）
# =====================================================================
s = add_slide(); header(s, "ANTI-HALLUCINATION · 检索侧", "宽召回 + 证据门槛：召得全、滤得准")
callout(s, Inches(0.62), Inches(1.42), Inches(12.1), Inches(1.02), "问题起点",
        "纯向量检索在中文药名上漏召：问“发烧能不能吃对乙酰氨基酚”，含药名片段语义分仅 0.48，"
        "被 search() 默认阈值 0.6 拦在候选之外 —— 重排救不回没进候选的片段。",
        fill=RGBColor(0xFB, 0xEF, 0xE9), accent=WARN, body_size=11.5)
bullets(s, Inches(0.62), Inches(2.72), Inches(5.55), Inches(3.0), [
    "宽召回 _wide_consultation_search：绕过 search() 阈值直查 collection，百级体量全量取，换确定性召回。",
    "短语加成 _query_phrases：抽 ≥5 字实体短语（滤虚词），命中即 strong_hit 并置顶。",
    "证据门槛 REACT_MIN_SCORE=0.50：只留真正相关片段；strong_hit 精确命中豁免门槛。",
    "过滤 _filter_by_score：exact 必须传 set 而非生成器（避免耗尽导致判定失真）。",
], size=12.5, gap=9)
pic(s, "4_retrieval_guard.png", Inches(6.35), Inches(2.62), Inches(6.45))
callout(s, Inches(0.62), Inches(5.85), Inches(12.1), Inches(0.85), "处置链",
        "宽召回捞回候选 → 短语加成识别 strong_hit → 豁免证据门槛并置顶 → 引用池去重后进入句子级溯源。",
        fill=LBLUE, accent=BLUE, body_size=11.5)
footer(s, 17)


# =====================================================================
# 18. 防幻觉（二）引用侧
# =====================================================================
s = add_slide(); header(s, "ANTI-HALLUCINATION · 引用侧", "引用池净化与句子级溯源")
bullets(s, Inches(0.62), Inches(1.5), Inches(12.1), Inches(5.0), [
    "引用池只用 kind=kb：记忆（既往对话）与就诊记录属于“个人史”而非“医学出处”，禁止占 [n] 编号 —— 否则会把个人史伪装成有知识出处（错误归因，比不标来源更糟）。",
    "去重与重排：引用池按 doc_id 去重再排序（-strong_hit, -score），避免“一轮 12 条引用里 9 条是 3 个片段的重复”。编号从最对症的片段开始。",
    "句子级溯源 build_sentence_trace：把终答逐句拆开、绑定每句引用的 [n]；cited=False 的句子前端显示“无直接知识来源”，既不丢内容也让用户看清哪些是通用表述。",
    "回校验 strip_undefined_cites：剥离模型乱标的未定义编号（如 [9]），防止小模型编造不存在的来源。",
    "兜底绑定：模型漏写 [n] 时按字面重合度就近挂源，但标 inferred=True（明确不是模型显式引用，不算有依据），避免出现伪溯源。",
    "二级兜底：零命中 → REACT_NO_EVIDENCE_REPLY（知识库未覆盖，建议就医）；超范围 → REACT_OUT_OF_SCOPE_REPLY；二者都不调模型、不推测。",
], size=13, gap=8)
footer(s, 18)


# =====================================================================
# 19. 防幻觉（三）工程侧
# =====================================================================
s = add_slide(); header(s, "ANTI-HALLUCINATION · 工程侧", "加密、确定性日期与透明可解释")
cards = [
    ("数据加密", "会话历史 / 报告 / 就诊记录落库前经 AES-256-GCM 加密；读取方必须先 decrypt_record() 再 json.loads，否则静默返回空。数据不出域。", BLUE),
    ("日期铁律", "任何日期一律由系统时间 current_time_str() 提供，并用 _enforce_visit_date() 兜底替换 / 补行；禁止 LLM 生成日期（实测会输出 2023 等旧日期）。", TEAL),
    ("可解释透出", "thinking 事件展示思维链、step 事件展示推理时间线，前端 ReasoningSteps 组件可折叠查看；透明 Observation 让模型无法虚构上下文。", DEEP),
    ("安全护栏", "EMERGENCY_KEYWORDS（胸痛 / 呼吸困难 / 昏迷等）命中先前置就医提醒；紧急场景优先生命安全而非科普作答。", WARN),
]
cw, ch = Inches(5.95), Inches(2.35)
gx, gy = Inches(0.62), Inches(1.55)
spx, spy = Inches(0.2), Inches(0.22)
for i, (t, b, col) in enumerate(cards):
    r, c = divmod(i, 2)
    x = gx + c*(cw+spx); y = gy + r*(ch+spy)
    callout(s, x, y, cw, ch, t, b, fill=LBLUE2, accent=col)
footer(s, 19)


# =====================================================================
# 20. 多模态与生态接入（新增）
# =====================================================================
s = add_slide(); header(s, "CAPABILITIES · 多模态与生态", "语音 / 图片 / 微信生态的扩展接入")
mm = [
    ("语音问诊（ASR）", "ASR_ENABLED 开关默认关闭；whisper base 模型 CPU int8 推理，中文医疗问诊提示词预热提升识别准确率；音频上限 10MB，支持常见格式。", BLUE),
    ("图片识别（Vision）", "VISION_ENABLED 开关默认关闭；视觉模型解读检查报告 / 患处照片，VISION_MAX_TOKENS=1500 控制输出长度，图片上限 10MB。", TEAL),
    ("微信生态", "订阅消息推送复诊提醒、wxacode 小程序码分享、位置服务推荐附近医院 —— 问诊闭环延伸到诊前与诊后。", DEEP),
    ("按需开启", "多模态能力全部经环境变量开关控制，离线最小部署不引入额外依赖；能力扩展不侵入 RAG / Agent 主链路。", WARN),
]
cw, ch = Inches(5.95), Inches(2.35)
gx, gy = Inches(0.62), Inches(1.55)
spx, spy = Inches(0.2), Inches(0.22)
for i, (t, b, col) in enumerate(mm):
    r, c = divmod(i, 2)
    x = gx + c*(cw+spx); y = gy + r*(ch+spy)
    callout(s, x, y, cw, ch, t, b, fill=LBLUE2, accent=col)
footer(s, 20)


# =====================================================================
# 21. 局限性与演进路线（新增）
# =====================================================================
s = add_slide(); header(s, "ROADMAP · 局限与演进", "当前边界与下一步路线")
bullets(s, Inches(0.62), Inches(1.5), Inches(5.6), Inches(4.8), [
    "知识规模：当前 11 篇文档 / 104 块，覆盖常见病与常用药，专科深度（如罕见病）仍需扩充语料。",
    "模型能力：7B 指令模型对复杂鉴别诊断推理有限，深度分析报告已分流至 deepseek-r1 推理模型。",
    "部署形态：单机本地部署，面向个人 / 家庭场景，未做高并发与多租户。",
], size=13, gap=10, lead_color=WARN)
road = [
    ("知识扩充", "批量导入权威医学指南 / 药品说明书；增量更新脚本保证分钟级生效。", BLUE),
    ("检索增强", "混合检索（BM25 + 向量）与可插拔重排（rerank），进一步提升中文药名召回。", TEAL),
    ("模型蒸馏", "training/ 已预留 QLoRA 对话微调与中文 embedding 微调管线，将高频问诊语料蒸馏进本地模型。", DEEP),
]
ry = Inches(1.55)
for i, (t, b, col) in enumerate(road):
    y = ry + i*Inches(1.52)
    callout(s, Inches(6.55), y, Inches(6.2), Inches(1.38), t, b, fill=LBLUE2, accent=col, body_size=11.5)
callout(s, Inches(0.62), Inches(6.25), Inches(12.1), Inches(0.62), "演进原则",
        "架构先行：RAG / Agent / 防幻觉三层机制保持稳定，语料、模型、检索算法均可在机制内平滑替换升级。",
        fill=LBLUE, accent=BLUE, body_size=11.5)
footer(s, 21)


# =====================================================================
# 22. 总结 + 测试分布实测（图 5 + 测试表）
# =====================================================================
s = add_slide(); header(s, "SUMMARY · 总结与实测", "三层能力 + 回归测试实测总览")
bullets(s, Inches(0.62), Inches(1.42), Inches(12.1), Inches(1.85), [
    "RAG：主题切分 + 宽召回 + 证据门槛，检索“召得全、滤得准、溯源稳”。",
    "Agent：ReAct 受控循环 + 工具白名单 + 双闸降级，7B 小模型可控、可恢复、可解释。",
    "防幻觉：检索侧（宽召回 / 强命中豁免）、引用侧（kb 专用池 / 句子级溯源 / 回校验）、工程侧（加密 / 日期 / 透明）三层设闸。",
], size=13, gap=7)
pic(s, "5_test_suite.png", Inches(0.55), Inches(3.35), Inches(7.15))
add_table(s, Inches(8.0), Inches(3.5), Inches(4.75), [
    ["测试指标", "实测结果"],
    ["用例总数", "77"],
    ["通过", "77（100%）"],
    ["执行耗时", "10.77s"],
    ["覆盖模块", "8 个"],
    ["知识库", "11 篇 → 104 块"],
], col_w=[1.1, 1.4], row_h=Pt(32), font_size=11)
callout(s, Inches(0.62), Inches(6.45), Inches(12.1), Inches(0.62), "核心价值",
        "医学回答做到“有出处、可追溯、不编造”—— 本地离线可部署（Ollama + 本地 ChromaDB），数据不出域，可解释性对齐医患信任。",
        fill=LBLUE, accent=BLUE, body_size=11.5)
footer(s, 22)


OUT = r"C:\Users\lenovo\Desktop\医疗问诊智能体_核心机制与原理_扩充版.pptx"
prs.save(OUT)
print("SAVED:", OUT, "slides:", len(prs.slides._sldIdLst))
